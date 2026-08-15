"""Export service for generating document exports."""
import json
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional

from docx import Document as DocxDocument
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches
from sqlalchemy.orm import Session

from app.config import settings
from app.models import Document, Export, ExportFormat, ExportStatus, Page, Block
from app.repositories import DocumentRepository, PageRepository, BlockRepository
from app.core import logger, ExportError


class ExportService:
    """Service for exporting documents."""

    def __init__(self, db: Session):
        self.db = db
        self.doc_repo = DocumentRepository(db)
        self.page_repo = PageRepository(db)
        self.block_repo = BlockRepository(db)

    def create_export(self, document_id: int, format: ExportFormat) -> Export:
        """Create a new export record."""
        export = Export(
            document_id=document_id,
            format=format,
            status=ExportStatus.PENDING,
        )
        self.db.add(export)
        self.db.commit()
        self.db.refresh(export)
        return export

    def export_document(
        self,
        document_id: int,
        format: ExportFormat,
        text_mode: Optional[str] = None,
    ) -> Export:
        """
        Export a document to the specified format.

        Args:
            document_id: Document ID
            format: Export format
            text_mode: Text selection mode for DOCX ("ocr_raw" or "llm_polished")

        Returns:
            Export record
        """
        document = self.doc_repo.get_by_id(document_id)
        if not document:
            raise ExportError(f"Document {document_id} not found")

        export = self.create_export(document_id, format)

        try:
            export.status = ExportStatus.PROCESSING
            self.db.commit()

            exports_dir = settings.get_exports_path() / str(document_id)
            exports_dir.mkdir(parents=True, exist_ok=True)

            timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
            filename = f"{document.original_filename}_{timestamp}"

            if format == ExportFormat.JSON:
                file_path = self._export_json(document, exports_dir, filename)
            elif format == ExportFormat.MARKDOWN:
                file_path = self._export_markdown(document, exports_dir, filename)
            elif format == ExportFormat.TXT:
                file_path = self._export_txt(document, exports_dir, filename)
            elif format == ExportFormat.DOCX:
                file_path = self._export_docx(document, exports_dir, filename, text_mode)
            else:
                raise ExportError(f"Unsupported export format: {format}")

            export.file_path = str(file_path)
            export.file_size = file_path.stat().st_size
            export.status = ExportStatus.COMPLETED
            self.db.commit()

            logger.info(f"Export completed: {file_path}")

        except Exception as e:
            logger.error(f"Export failed: {e}")
            export.status = ExportStatus.FAILED
            self.db.commit()
            raise ExportError(str(e))

        return export

    def _export_json(self, document: Document, output_dir: Path, filename: str) -> Path:
        """Export document to JSON format."""
        pages = self.page_repo.get_by_document(document.id)

        data = {
            "document": {
                "id": document.id,
                "filename": document.original_filename,
                "page_count": document.page_count,
                "created_at": document.created_at.isoformat(),
            },
            "pages": [],
        }

        for page in pages:
            blocks = self.block_repo.get_by_page(page.id)
            page_data = {
                "page_number": page.page_number,
                "markdown": page.markdown_content,
                "blocks": [
                    {
                        "index": block.block_index,
                        "label": block.block_label,
                        "bbox": block.bbox,
                        "ocr_raw": block.ocr_text_raw,
                        "ocr_corrected": block.ocr_text_corrected,
                        "modern_english": block.modern_english,
                        "translation_zh": block.translation_zh,
                    }
                    for block in blocks
                ],
            }
            data["pages"].append(page_data)

        file_path = output_dir / f"{filename}.json"
        file_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        return file_path

    def _export_markdown(self, document: Document, output_dir: Path, filename: str) -> Path:
        """Export document to Markdown format."""
        pages = self.page_repo.get_by_document(document.id)

        lines = [
            f"# {document.original_filename}",
            "",
            f"Exported: {datetime.utcnow().isoformat()}",
            "",
            "---",
            "",
        ]

        for page in pages:
            lines.append(f"## Page {page.page_number}")
            lines.append("")

            blocks = self.block_repo.get_by_page(page.id)
            for block in blocks:
                if block.ocr_text_raw:
                    lines.append(f"### Block {block.block_index + 1} ({block.block_label})")
                    lines.append("")
                    lines.append("**Original:**")
                    lines.append(f"> {block.ocr_text_raw}")
                    lines.append("")

                    if block.ocr_text_corrected:
                        lines.append("**Corrected:**")
                        lines.append(f"> {block.ocr_text_corrected}")
                        lines.append("")

                    if block.modern_english:
                        lines.append("**Modern English:**")
                        lines.append(f"> {block.modern_english}")
                        lines.append("")

                    if block.translation_zh:
                        lines.append("**中文翻译:**")
                        lines.append(f"> {block.translation_zh}")
                        lines.append("")

            lines.append("---")
            lines.append("")

        file_path = output_dir / f"{filename}.md"
        file_path.write_text("\n".join(lines), encoding="utf-8")
        return file_path

    def _export_txt(self, document: Document, output_dir: Path, filename: str) -> Path:
        """Export document to plain text format."""
        pages = self.page_repo.get_by_document(document.id)

        lines = [
            document.original_filename,
            "=" * len(document.original_filename),
            "",
        ]

        for page in pages:
            lines.append(f"--- Page {page.page_number} ---")
            lines.append("")

            blocks = self.block_repo.get_by_page(page.id)
            for block in blocks:
                text = self._get_best_block_text(block)
                if text:
                    lines.append(text)
                    lines.append("")

        file_path = output_dir / f"{filename}.txt"
        file_path.write_text("\n".join(lines), encoding="utf-8")
        return file_path

    def _export_docx(
        self,
        document: Document,
        output_dir: Path,
        filename: str,
        text_mode: Optional[str] = None,
    ) -> Path:
        """Export document to DOCX format with image+text page pairs.

        Args:
            document: Document model
            output_dir: Output directory
            filename: Base filename
            text_mode: "ocr_raw" for raw OCR text, "llm_polished" (default) for best available
        """
        use_raw = text_mode == "ocr_raw"
        pages = self.page_repo.get_by_document(document.id)
        doc = DocxDocument()
        doc.core_properties.title = document.original_filename
        mode_label = "OCR原文" if use_raw else "OCR纠错"
        doc.core_properties.subject = f"OCR comparison export ({mode_label})"
        doc.core_properties.comments = "Generated from persisted OCR and GPT-polished block data."

        for index, page in enumerate(pages):
            # --- Image page ---
            if index > 0:
                doc.add_section(WD_SECTION_START.NEW_PAGE)

            self._configure_docx_page(doc, page)
            self._append_image_page(doc, page)

            # --- Text page ---
            doc.add_section(WD_SECTION_START.NEW_PAGE)
            self._configure_docx_page(doc, page)

            blocks = self._get_sorted_blocks(self.block_repo.get_by_page(page.id), page)
            self._append_page_header(doc, page, len(blocks), use_raw)
            self._append_page_blocks(doc, page, blocks, use_raw)

        suffix = "_ocr" if use_raw else "_polished"
        file_path = output_dir / f"{filename}{suffix}.docx"
        doc.save(file_path)
        return file_path

    def _append_image_page(self, doc: DocxDocument, page: Page) -> None:
        """Insert the original page image centered on a DOCX page."""
        image_path = Path(page.image_path) if page.image_path else None
        if not image_path or not image_path.exists():
            # No image available — add a placeholder paragraph
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"Page {page.page_number} — 原图不可用")
            run.italic = True
            run.font.size = Pt(10)
            return

        # Header label
        header = doc.add_paragraph()
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header.paragraph_format.space_after = Pt(6)
        header_run = header.add_run(f"Page {page.page_number} — 原文")
        header_run.bold = True
        header_run.font.size = Pt(11)

        # Insert image scaled to fit the usable page width
        section = doc.sections[-1]
        usable_width = section.page_width - section.left_margin - section.right_margin
        usable_height = section.page_height - section.top_margin - section.bottom_margin - Inches(0.5)

        # Calculate image aspect ratio to fit within usable area
        img_width = max(page.width or 1, 1)
        img_height = max(page.height or 1, 1)
        scale_w = float(usable_width) / img_width
        scale_h = float(usable_height) / img_height
        scale = min(scale_w, scale_h)
        final_width = int(img_width * scale)

        try:
            doc.add_picture(str(image_path), width=final_width)
            # Center the image
            last_paragraph = doc.paragraphs[-1]
            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as e:
            logger.warning(f"Failed to add page image: {e}")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"Page {page.page_number} — 图片加载失败")
            run.italic = True
            run.font.size = Pt(10)

    def _configure_docx_page(self, doc: DocxDocument, page: Page) -> None:
        """Configure DOCX page size to roughly match source page geometry."""
        section = doc.sections[-1]
        width = max(page.width or 1000, 1)
        height = max(page.height or 1400, 1)
        longest_edge = max(width, height)
        scale = 10.0 / longest_edge
        page_width_in = max(width * scale, 5.5)
        page_height_in = max(height * scale, 7.0)
        section.page_width = Inches(page_width_in)
        section.page_height = Inches(page_height_in)
        margin = Inches(0.35)
        section.top_margin = margin
        section.bottom_margin = margin
        section.left_margin = margin
        section.right_margin = margin

    def _append_page_header(
        self, doc: DocxDocument, page: Page, block_count: int, use_raw: bool = False
    ) -> None:
        """Append a compact page header for comparison-friendly exports."""
        mode_label = "OCR原文" if use_raw else "OCR纠错"
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(2)
        title_run = title.add_run(f"Page {page.page_number} — {mode_label}")
        title_run.bold = True
        title_run.font.size = Pt(11)

        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.paragraph_format.space_after = Pt(6)
        meta_run = meta.add_run(
            f"{page.width or '-'} × {page.height or '-'} · {block_count} OCR blocks"
        )
        meta_run.italic = True
        meta_run.font.size = Pt(8)

    def _append_page_blocks(
        self, doc: DocxDocument, page: Page, blocks: Iterable[Block], use_raw: bool = False
    ) -> None:
        """Append ordered block content with approximate coordinate-based layout."""
        previous_bottom = 0.0
        page_height = float(max(page.height or 1, 1))
        page_width = float(max(page.width or 1, 1))
        content_width_points = self._get_content_width_points(doc.sections[-1])

        for block in blocks:
            text = self._get_block_text_by_mode(block, use_raw)
            if not text:
                continue

            bbox = self._normalize_bbox(block.bbox)
            top = bbox[1]
            bottom = bbox[3]
            left = bbox[0]
            gap = max(top - previous_bottom, 0.0)
            spacing_before = min((gap / page_height) * 720.0, 48.0)
            left_indent = min((left / page_width) * content_width_points, content_width_points * 0.55)

            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(spacing_before)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.left_indent = Pt(left_indent)

            text_run = paragraph.add_run(text)
            text_run.font.size = Pt(10.5)

            previous_bottom = max(previous_bottom, bottom)

    def _get_sorted_blocks(self, blocks: Iterable[Block], page: Page) -> list[Block]:
        """Sort blocks by visual top/left order for export."""
        page_height = float(max(page.height or 1, 1))

        def sort_key(block: Block) -> tuple[float, float, int]:
            bbox = self._normalize_bbox(block.bbox)
            row_bucket = round((bbox[1] / page_height) * 1000)
            return row_bucket, bbox[0], block.block_index

        return sorted(blocks, key=sort_key)

    def _normalize_bbox(self, bbox: object) -> list[float]:
        """Normalize stored bbox values for layout calculations."""
        if not isinstance(bbox, list) or len(bbox) < 4:
            return [0.0, 0.0, 0.0, 0.0]
        try:
            return [float(bbox[0]), float(bbox[1]), float(bbox[2]), float(bbox[3])]
        except (TypeError, ValueError):
            return [0.0, 0.0, 0.0, 0.0]

    def _get_content_width_points(self, section) -> float:
        """Return usable section width in points."""
        return max(float(section.page_width - section.left_margin - section.right_margin) / 12700.0, 1.0)

    def _get_block_text_by_mode(self, block: Block, use_raw: bool) -> str:
        """Return block text according to the selected export mode.

        Args:
            block: Block model
            use_raw: If True, return only raw OCR text; otherwise corrected text

        Returns:
            Selected text content
        """
        if use_raw:
            return block.ocr_text_raw or ""
        return block.ocr_text_corrected or block.ocr_text_raw or ""

    def _get_best_block_text(self, block: Block) -> str:
        """Return the best persisted text for a block."""
        return (
            block.translation_zh
            or block.modern_english
            or block.ocr_text_corrected
            or block.ocr_text_raw
            or ""
        )

    def get_exports(self, document_id: int) -> list[Export]:
        """Get all exports for a document."""
        return (
            self.db.query(Export)
            .filter(Export.document_id == document_id)
            .order_by(Export.created_at.desc())
            .all()
        )
